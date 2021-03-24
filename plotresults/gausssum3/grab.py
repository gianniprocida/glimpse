import pandas as pd
import docx
c=8065.73 
osc,energy=[],[] 
i,j=[],[]
contr=[] 
k,e=[],[]
exc=['E1','E2','E3','E4','E5','E6']
searchfile=open('UVData.txt','r') 
for line in searchfile: 
        if 'HOMO' in line: 
            next(searchfile) 
            for line in searchfile: 
                osc.append(line.split()[3]) 
                i.append(line.split()[5]) 
                j.append(line.split()[6])
                k.append(line.split()[7])
                e.append(line.split()[8])
                energy.append(float(line.split()[1])/c)
                 
first=[a+b for a,b in zip(i,j)]
second=[a+b for a,b in zip(k,e)]
comp=[a+b for a,b in zip(first,second)]

df=pd.DataFrame(dict(Excitation=exc, Energy=energy,OS=osc,Composition=comp))


df=df.round(2)
doc = docx.Document('test.docx')


t = doc.add_table(df.shape[0]+1, df.shape[1])

# add the header rows.
for j in range(df.shape[-1]):
    t.cell(0,j).text = df.columns[j]

# add the rest of the data frame
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        t.cell(i+1,j).text = str(df.values[i,j])

# save the doc
doc.save('test.docx')



